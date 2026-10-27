"""Reconcile source catalogues against the original 549 and new fiche manifests."""
import hashlib
import json
import sys
from collections import defaultdict
from pathlib import Path

from pptx import Presentation


def digest(path):
    sha = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()


def sady_slide_equivalences(original, inventory):
    """Match split SADY decks to original slides by unique photo or exact text."""
    rows = [row for row in original["registros"] if row.get("proveedor") == "SADY"]
    if not rows:
        return {}
    original_deck = Path(rows[0]["archivo_fuente"])
    if not original_deck.is_file():
        return {}
    slides_by_photo = defaultdict(set)
    slides_by_text = defaultdict(set)
    for number, slide in enumerate(Presentation(original_deck).slides, 1):
        slide_text = " ".join(" ".join(shape.text.split()) for shape in slide.shapes if shape.has_text_frame)
        slides_by_text[slide_text].add(number)
        for shape in slide.shapes:
            if shape.shape_type == 13 and len(shape.image.blob) > 20_000:
                slides_by_photo[hashlib.sha256(shape.image.blob).hexdigest()].add(number)
    code_by_slide = {row["pagina_diapositiva"]: row["codigo_hercas"] for row in rows}
    matches = {}
    for source in inventory["sources"]:
        path = Path(source["path"])
        if source.get("error") or path.suffix.lower() != ".pptx" or "6. SADY" not in source["relative"]:
            continue
        try:
            slide = Presentation(path).slides[0]
            candidates = set()
            for shape in slide.shapes:
                if shape.shape_type != 13 or len(shape.image.blob) <= 20_000:
                    continue
                candidate = slides_by_photo.get(hashlib.sha256(shape.image.blob).hexdigest(), set())
                if len(candidate) == 1:
                    candidates.update(candidate)
            if len(candidates) == 1:
                number = candidates.pop()
                if number in code_by_slide:
                    matches[source["path"]] = (str(original_deck), code_by_slide[number], number, "foto")
            elif not candidates:
                slide_text = " ".join(" ".join(shape.text.split()) for shape in slide.shapes if shape.has_text_frame)
                text_candidates = slides_by_text.get(slide_text, set())
                if len(text_candidates) == 1:
                    number = next(iter(text_candidates))
                    if number in code_by_slide:
                        matches[source["path"]] = (str(original_deck), code_by_slide[number], number, "texto")
        except (OSError, ValueError, IndexError):
            continue
    return matches


def deck_content_digest(path):
    """Hash slide text and embedded pictures, ignoring Office package metadata."""
    sha = hashlib.sha256()
    for slide in Presentation(path).slides:
        sha.update(b"SLIDE\0")
        for shape in slide.shapes:
            if shape.has_text_frame:
                sha.update(shape.text.encode("utf-8"))
            if shape.shape_type == 13:
                sha.update(hashlib.sha256(shape.image.blob).digest())
    return sha.hexdigest()


def main(root):
    original = json.loads((root / "avance.json").read_text(encoding="utf-8"))
    inventory = json.loads((root / "_nuevos_medios" / "inventario_fuentes.json").read_text(encoding="utf-8"))
    old_sources = {}
    for name in {row["archivo_fuente"] for row in original["registros"]}:
        file = Path(name)
        if file.is_file():
            old_sources[digest(file)] = str(file)
    sady_matches = sady_slide_equivalences(original, inventory)
    fiches = defaultdict(list)
    for file in root.rglob("origen_y_control.json"):
        try:
            data = json.loads(file.read_text(encoding="utf-8"))
            if Path(data.get("fuente", "")).is_file():
                fiches[digest(Path(data["fuente"]))].append(data["codigo"])
        except (OSError, ValueError):
            continue
    first_for_sha = {}
    first_for_content = {}
    entries = []
    for source in inventory["sources"]:
        sha = source["sha256"]
        content_sha = None
        if not source.get("error") and source.get("pages", 0) > 1 and source.get("extension") == ".pptx":
            content_sha = deck_content_digest(source["path"])
        if source.get("error"):
            status = "requiere_conversion_o_reparacion"
        elif sha in first_for_sha:
            status = "copia_duplicada"
        elif content_sha and content_sha in first_for_content:
            status = "copia_duplicada_contenido"
        elif sha in old_sources:
            status = "archivo_ya_en_549"
        elif source["path"] in sady_matches:
            status = "medio_ya_en_549"
        elif sha in fiches:
            status = "en_produccion"
        else:
            status = "pendiente_desglose"
        entries.append({
            "archivo": source["path"], "relativo": source["relative"],
            "area": source["area"], "grupo": source["source_group"],
            "sha256": sha, "paginas": source.get("pages"), "estado": status,
            "duplicado_de": first_for_sha.get(sha) or (first_for_content.get(content_sha) if content_sha else None),
            "archivo_en_549": old_sources.get(sha) or sady_matches.get(source["path"], (None, None, None, None))[0],
            "codigo_en_549": sady_matches.get(source["path"], (None, None, None, None))[1],
            "diapositiva_en_549": sady_matches.get(source["path"], (None, None, None, None))[2],
            "fichas_producidas": sorted(set(fiches.get(sha, []))),
            "observacion": source.get("error", "") or ("Foto o texto idéntico a la diapositiva original" if source["path"] in sady_matches else ""),
        })
        first_for_sha.setdefault(sha, source["path"])
        if content_sha:
            first_for_content.setdefault(content_sha, source["path"])
    report = {
        "raiz_fuentes": inventory["root"],
        "total_archivos": len(entries),
        "archivos_unicos": len(first_for_sha),
        "total_fichas_nuevas": sum(len(x) for x in fiches.values()),
        "fuentes": entries,
    }
    output = root / "_nuevos_medios" / "control_fuentes.json"
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    from collections import Counter
    print(json.dumps({"output": str(output), "estados": Counter(x["estado"] for x in entries), "fichas": report["total_fichas_nuevas"]}, ensure_ascii=False))


if __name__ == "__main__":
    main(Path(sys.argv[1]))
