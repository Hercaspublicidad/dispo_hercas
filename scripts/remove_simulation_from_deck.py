import argparse
from pathlib import Path
from pptx import Presentation


def clean(path: Path):
    deck = Presentation(path)
    slide = deck.slides[0]
    removed = 0
    for shape in list(slide.shapes):
        text = shape.text if shape.has_text_frame else ''
        if shape.name in {'simulacion', 'franja inferior foto', 'acento inferior foto'} or 'SIMULACIÓN DE PAUTA HERCAS' in text.upper():
            shape._element.getparent().remove(shape._element)
            removed += 1
    if not any(shape.name == 'precio editable' and shape.has_text_frame for shape in slide.shapes):
        raise ValueError('No hay precio editable en la ficha')
    if removed:
        tmp = path.with_suffix('.cleaning.pptx')
        deck.save(tmp)
        tmp.replace(path)
    return removed


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('pptx', type=Path)
    args = parser.parse_args()
    print(f'Elementos retirados: {clean(args.pptx)}')
