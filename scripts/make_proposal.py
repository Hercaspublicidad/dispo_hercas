import argparse
import json
from collections import OrderedDict
from datetime import date
from copy import deepcopy
from io import BytesIO
from math import cos, sin, pi
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

NAVY = RGBColor(6, 25, 51)
YELLOW = RGBColor(255, 207, 20)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(198, 210, 228)

def add_text(slide, text, left, top, width, height, size, color, bold=False, name=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    if name:
        box.name = name
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    paragraph = tf.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box

def remove_first_slide(deck):
    slide_id = deck.slides._sldIdLst[0]
    deck.part.drop_rel(slide_id.rId)
    deck.slides._sldIdLst.remove(slide_id)

def blank_slide(deck):
    layout = next((x for x in deck.slide_layouts if "blank" in x.name.lower() or "blanco" in x.name.lower()), deck.slide_layouts[0])
    return deck.slides.add_slide(layout)

def add_cover(deck, root, count, city_count):
    slide = blank_slide(deck)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    model = root / "modelo_hercas_portada.jpg"
    if model.is_file():
        with Image.open(model) as im:
            right = im.crop((int(im.width * .50), 0, im.width, im.height))
            blob = BytesIO()
            right.convert("RGB").save(blob, format="JPEG", quality=90)
            blob.seek(0)
            slide.shapes.add_picture(blob, Inches(6.2), 0, width=deck.slide_width-Inches(6.2), height=deck.slide_height)
    add_text(slide, "HERCAS", Inches(.8), Inches(.55), Inches(4.7), Inches(.5), 23, WHITE, True, "marca")
    add_text(slide, "VALLAS\nDISPONIBLES", Inches(.8), Inches(1.8), Inches(5.4), Inches(2.7), 43, YELLOW, True, "titulo portada")
    add_text(slide, f"{count} vallas seleccionadas  |  {city_count} ciudades", Inches(.8), Inches(5.5), Inches(5.2), Inches(.45), 17, WHITE, False, "resumen")
    add_text(slide, "Disponibilidad sujeta a confirmación antes de contratar.", Inches(.8), Inches(6.65), Inches(5.3), Inches(.4), 12, MUTED, False, "vigencia")

def alpha_fill(shape, rgb, opacity):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(int(opacity * 100000)))
    shape._element.spPr.solidFill.srgbClr.append(alpha)
    shape.line.fill.background()

def add_outline_text(slide, text, left, top, width, height, size):
    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = "zona contorno"
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = False
    rpr = run._r.get_or_add_rPr()
    # PowerPoint requires the outline before noFill/latin in run properties.
    line = OxmlElement("a:ln")
    line.set("w", "22860")
    solid = OxmlElement("a:solidFill")
    color = OxmlElement("a:srgbClr")
    color.set("val", "FFFFFF")
    solid.append(color)
    line.append(solid)
    rpr.insert(0, line)
    rpr.insert(1, OxmlElement("a:noFill"))
    return box

def add_city(deck, city, count, photo_path):
    if not photo_path or not Path(photo_path).is_file():
        raise ValueError(f"Falta foto de fondo para {city}")
    slide = blank_slide(deck)
    with Image.open(photo_path) as im:
        ratio = im.width / im.height
    frame = deck.slide_width / deck.slide_height
    picture = slide.shapes.add_picture(str(photo_path), 0, 0, width=deck.slide_width, height=deck.slide_height)
    picture.name = "foto ciudad"
    if ratio > frame:
        crop = (1 - frame / ratio) / 2
        picture.crop_left = picture.crop_right = crop
    else:
        crop = (1 - ratio / frame) / 2
        picture.crop_top = picture.crop_bottom = crop
    overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, deck.slide_width, deck.slide_height)
    overlay.name = "filtro azul"
    alpha_fill(overlay, RGBColor(0, 20, 70), .78)
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.5), Inches(2.03), Inches(5.9), Inches(3.55))
    circle.name = "halo azul"
    alpha_fill(circle, RGBColor(0, 17, 62), .95)
    semi = slide.shapes.build_freeform(0, 0, scale=Inches(.9) / 1000)
    semi.add_line_segments([(1000, 0)] + [(500 + 500*cos(pi*i/32), 500*sin(pi*i/32)) for i in range(1, 33)], close=True)
    accent = semi.convert_to_shape(Inches(3.12), Inches(2.68))
    accent.name = "semicírculo amarillo"
    accent.fill.solid()
    accent.fill.fore_color.rgb = YELLOW
    accent.line.fill.background()
    add_outline_text(slide, "ZONA", Inches(4.18), Inches(2.45), Inches(6.2), Inches(.9), 57)
    city_display = city.replace("Bogot�", "Bogotá").replace("Medell�n", "Medellín").replace("Itag��", "Itagüí").replace("Jer�nimo", "Jerónimo")
    display = city_display.upper()
    size = 60 if len(display) <= 13 else (48 if len(display) <= 20 else 40)
    if len(display) > 22 and " Y " in display:
        display = display.replace(" Y ", " Y\n", 1)
    add_text(slide, display, Inches(4.18), Inches(3.35), Inches(8.2), Inches(2.0), size, WHITE, True, "ciudad")

def image_stream(shape, compress):
    blob = shape.image.blob
    if not compress:
        return BytesIO(blob)
    with Image.open(BytesIO(blob)) as im:
        converted = im.convert("RGB")
        output = BytesIO()
        converted.save(output, format="JPEG", quality=90, optimize=True)
        output.seek(0)
        return output

def copy_slide(deck, source):
    target = blank_slide(deck)
    for i, shape in enumerate(source.shapes):
        if shape.name in ("simulacion", "franja inferior foto") or (shape.has_text_frame and "SIMULACIÓN DE PAUTA HERCAS" in shape.text.upper()):
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            new = target.shapes.add_picture(image_stream(shape, i == 4), shape.left, shape.top, width=shape.width, height=shape.height)
            new.crop_left, new.crop_right = shape.crop_left, shape.crop_right
            new.crop_top, new.crop_bottom = shape.crop_top, shape.crop_bottom
            new.rotation = shape.rotation
            new.name = shape.name
        else:
            target.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    return target

def create(root, codes, output):
    manifest = json.loads((root / "avance.json").read_text(encoding="utf-8"))
    by_code = {r["codigo_hercas"]: r for r in manifest["registros"]}
    override_file = root / "catalogo_disponibilidad.json"
    overrides = json.loads(override_file.read_text(encoding="utf-8")) if override_file.exists() else {}
    if not codes or len(codes) > 550 or len(set(codes)) != len(codes):
        raise ValueError("Seleccione entre 1 y 550 códigos únicos.")
    records = []
    for code in codes:
        r = by_code.get(code)
        if not r or r.get("estado") != "completo" or not r.get("ficha_editable") or not Path(r["ficha_editable"]).is_file():
            raise ValueError(f"Ficha no disponible: {code}")
        override = overrides.get(code, {})
        availability = str(override.get("disponibilidad") or r.get("disponibilidad_fuente") or "").lower()
        source_date = override.get("fecha") or r.get("fecha_fuente")
        try:
            recent = abs((date.today() - date.fromisoformat(str(source_date)[:10])).days) <= 14
        except (TypeError, ValueError):
            recent = False
        if recent and "no disponible" in availability:
            raise ValueError(f"Valla con indisponibilidad vigente: {code}")
        records.append(r)
    groups = OrderedDict()
    for r in sorted(records, key=lambda x: (str(x.get("ciudad") or "").casefold(), x["codigo_hercas"])):
        groups.setdefault(r.get("ciudad") or "Ciudad por confirmar", []).append(r)
    first = Presentation(records[0]["ficha_editable"])
    width, height = first.slide_width, first.slide_height
    remove_first_slide(first)
    add_cover(first, root, len(records), len(groups))
    ordered = []
    photo_codes = json.loads((Path(__file__).resolve().parents[1] / "city_backgrounds.json").read_text(encoding="utf-8"))
    for city, items in groups.items():
        background_record = by_code.get(photo_codes.get(city)) or items[0]
        add_city(first, city, len(items), background_record.get("foto_HERCAS"))
        ordered.append({"type": "ciudad", "city": city})
        for r in items:
            source = Presentation(r["ficha_editable"])
            if source.slide_width != width or source.slide_height != height:
                raise ValueError(f"Tamaño de ficha incompatible: {r['codigo_hercas']}")
            copy_slide(first, source.slides[0])
            ordered.append({"type": "valla", "city": city, "code": r["codigo_hercas"]})
    output.parent.mkdir(parents=True, exist_ok=True)
    first.save(output)
    check = Presentation(output)
    if len(check.slides) != 1 + len(groups) + len(records):
        raise ValueError("El número de diapositivas no coincide.")
    for i, item in enumerate(ordered, 1):
        if item["type"] == "valla":
            shapes = {s.name: s for s in check.slides[i].shapes}
            if item["code"] not in shapes["titulo"].text or not shapes["precio editable"].has_text_frame:
                raise ValueError(f"Ficha no editable o código incorrecto: {item['code']}")
    return {"slides": len(check.slides), "vallas": len(records), "ciudades": list(groups), "bytes": output.stat().st_size}

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--codes", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    result = create(Path(args.root), json.loads(args.codes), Path(args.output))
    print(json.dumps(result, ensure_ascii=False))

if __name__ == "__main__":
    main()
