"""Read-only inventory of the VALLAS, CC, AEROPUERTO source catalogues."""
import hashlib
import json
import sys
from pathlib import Path

from pypdf import PdfReader
from pptx import Presentation


def fingerprint(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def describe(path, root):
    relative = path.relative_to(root)
    result = {
        "path": str(path), "relative": str(relative), "area": relative.parts[0],
        "source_group": relative.parts[1] if len(relative.parts) > 2 else "",
        "extension": ".pptx" if path.name.lower() == ".pptx" else path.suffix.lower(), "bytes": path.stat().st_size,
        "sha256": fingerprint(path), "pages": None, "samples": [], "images": None,
    }
    try:
        if result["extension"] == ".pdf":
            doc = PdfReader(str(path))
            result["pages"] = len(doc.pages)
            result["samples"] = [
                {"page": i + 1, "text": " ".join((page.extract_text() or "").split())[:700]}
                for i, page in enumerate(doc.pages[:min(4, len(doc.pages))])
            ]
        elif result["extension"] == ".pptx":
            deck = Presentation(str(path))
            result["pages"] = len(deck.slides)
            result["images"] = sum(
                shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes
            )
            result["samples"] = [
                {"page": i + 1, "text": " ".join(
                    shape.text for shape in slide.shapes if shape.has_text_frame
                ).strip()[:700]}
                for i, slide in enumerate(list(deck.slides)[:4])
            ]
        elif result["extension"] == ".ppt":
            result["error"] = "PowerPoint 97-2003: convert to PPTX before extraction"
        else:
            result["error"] = "Unknown format"
    except Exception as error:
        result["error"] = str(error)
    return result


def main():
    root = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    sources = sorted(path for path in root.rglob("*") if path.is_file())
    inventory = {"root": str(root), "sources": [describe(path, root) for path in sources]}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "sources": len(sources),
        "pages": sum(source["pages"] or 0 for source in inventory["sources"]),
        "errors": [source["relative"] for source in inventory["sources"] if source.get("error")],
        "output": str(output),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
