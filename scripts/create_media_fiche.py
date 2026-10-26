"""Build a HERCAS PowerPoint fiche and a lightweight preview from a verified montage."""
import argparse
import json
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ROOT = Path(r"D:\Aliados exteriores\Fichas_HERCAS_2026")
TEMPLATE = ROOT / "METRO VIAS" / "Barranquilla" / "MV-BAQ-001" / "MV-BAQ-001_ficha_HERCAS_editable.pptx"


def set_text(shape, value):
    frame = shape.text_frame
    for paragraph in frame.paragraphs:
        for run in paragraph.runs:
            run.text = ""
    if frame.paragraphs[0].runs:
        frame.paragraphs[0].runs[0].text = str(value)
    else:
        frame.paragraphs[0].text = str(value)


def replace_photo(slide, image_path):
    old = slide.shapes[4]
    if old.shape_type != MSO_SHAPE_TYPE.PICTURE:
        raise ValueError("La plantilla cambió: no se encontró la foto principal")
    parent = old._element.getparent()
    position = parent.index(old._element)
    new = slide.shapes.add_picture(str(image_path), old.left, old.top, width=old.width, height=old.height)
    with Image.open(image_path) as photo:
        ratio = photo.width / photo.height
    frame = old.width / old.height
    if ratio > frame:
        new.crop_left = new.crop_right = (1 - frame / ratio) / 2
    else:
        new.crop_top = new.crop_bottom = (1 - ratio / frame) / 2
    parent.remove(new._element)
    parent.insert(position, new._element)
    parent.remove(old._element)


def preview(photo_path, destination, record):
    canvas = Image.new("RGB", (1600, 900), "#0b2037")
    with Image.open(photo_path) as source:
        image = source.convert("RGB")
        image.thumbnail((1512, 570))
        canvas.paste(image, ((1600 - image.width) // 2, 172 + (570 - image.height) // 2))
    draw = ImageDraw.Draw(canvas)
    font_path = r"C:\Windows\Fonts\arial.ttf"
    bold_path = r"C:\Windows\Fonts\arialbd.ttf"
    title = ImageFont.truetype(bold_path, 52)
    label = ImageFont.truetype(font_path, 28)
    small = ImageFont.truetype(font_path, 24)
    draw.text((50, 40), record["codigo"] + " · " + record["nombre"], font=title, fill="white")
    draw.text((50, 112), record["ciudad"] + " · " + record["entorno"].replace("_", " ").title(), font=label, fill="#f7d548")
    draw.text((50, 770), record["ubicacion"], font=label, fill="white")
    draw.text((50, 820), record["formato"] + " · " + record.get("medidas", "Medidas por confirmar"), font=small, fill="#d6e0eb")
    draw.text((1150, 820), record["precio_texto"], font=small, fill="#f7d548")
    canvas.save(destination, optimize=True)


def create(record, image_path):
    code = record["codigo"]
    folder = ROOT / record["grupo_carpeta"] / record["ciudad"] / code
    folder.mkdir(parents=True, exist_ok=True)
    photo = folder / f"{code}_foto_HERCAS.png"
    fiche = folder / f"{code}_ficha_HERCAS_editable.pptx"
    thumb = folder / f"{code}_ficha_HERCAS.png"
    if not photo.exists():
        shutil.copyfile(image_path, photo)
    deck = Presentation(TEMPLATE)
    slide = deck.slides[0]
    for shape in list(slide.shapes):
        if shape.name in ("simulacion", "franja inferior foto") or (shape.has_text_frame and "SIMULACIÓN DE PAUTA HERCAS" in shape.text.upper()):
            shape._element.getparent().remove(shape._element)
    replace_photo(slide, photo)
    shapes = {shape.name: shape for shape in slide.shapes}
    values = {
        "titulo": code + " · " + record["nombre"],
        "subtitulo": record["grupo_carpeta"] + " · " + record["ciudad"],
        "ubicacion": record["ubicacion"],
        "sentido": "Sentido: " + record.get("sentido", "Por confirmar"),
        "formato": record["formato"] + " · " + record.get("medidas", "Medidas por confirmar"),
        "entorno": record["entorno"].replace("_", " ").title(),
        "precio editable": record["precio_texto"],
        "condicion precio": record.get("condicion_precio", "+ IVA / mes"),
    }
    for name, value in values.items():
        set_text(shapes[name], value)
    deck.save(fiche)
    preview(photo, thumb, record)
    check = Presentation(fiche)
    price = next(shape for shape in check.slides[0].shapes if shape.name == "precio editable")
    if not price.has_text_frame or price.text != record["precio_texto"]:
        raise AssertionError("El precio no quedó editable")
    manifest = {**record, "foto": str(photo), "ficha": str(fiche), "vista_previa": str(thumb)}
    (folder / "origen_y_control.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"codigo": code, "foto": str(photo), "ficha": str(fiche), "vista_previa": str(thumb)}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("record")
    parser.add_argument("image")
    args = parser.parse_args()
    create(json.loads(Path(args.record).read_text(encoding="utf-8")), Path(args.image))
